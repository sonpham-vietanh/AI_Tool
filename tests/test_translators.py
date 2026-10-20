import asyncio
import zipfile

import docx
import pytest

from backend.translators import translate_docx_file, translate_epub_file, translate_txt_file


def run(coro):
    return asyncio.run(coro)


# ------------------------------------------------------------------- TXT
TXT = (
    "First paragraph here.\n\n"
    "Second paragraph here.\n\n"
    "```python\nprint('keep me')\n```\n\n"
    "---\n\n"
    "12345\n\n"
    "https://example.com/x\n\n"
    "Last paragraph here.\n"
)


def test_txt_end_to_end(tmp_path, make_client):
    src, dst = tmp_path / "in.txt", tmp_path / "out" / "out.txt"
    src.write_text(TXT, encoding="utf-8")
    cl = make_client()
    res = run(translate_txt_file(str(src), str(dst), cl))
    assert res["status"] == "completed" and res["total_items"] == 3
    out = dst.read_text(encoding="utf-8")
    for p in ("First paragraph here.", "Second paragraph here.", "Last paragraph here."):
        assert f"VI({p})" in out
    # protected/passthrough content untouched
    assert "print('keep me')" in out and "VI(print" not in out
    assert "12345" in out and "https://example.com/x" in out and "VI(12345" not in out
    # paragraph structure preserved
    assert out.count("\n\n") == TXT.count("\n\n")


def test_txt_bilingual(tmp_path, make_client):
    src, dst = tmp_path / "in.txt", tmp_path / "out.txt"
    src.write_text("Hello world.\n\nGoodbye world.\n", encoding="utf-8")
    run(translate_txt_file(str(src), str(dst), make_client(), bilingual=True))
    out = dst.read_text(encoding="utf-8")
    assert "Hello world.\n\nVI(Hello world.)" in out


def test_txt_no_translatable_copies_file(tmp_path, make_client):
    src, dst = tmp_path / "in.txt", tmp_path / "out.txt"
    src.write_text("123\n\n---\n", encoding="utf-8")
    cl = make_client()
    res = run(translate_txt_file(str(src), str(dst), cl))
    assert res["total_items"] == 0 and dst.read_text(encoding="utf-8") == "123\n\n---\n"
    assert cl.calls == []


def test_txt_progress_and_cancel(tmp_path, make_client):
    src = tmp_path / "in.txt"
    src.write_text("\n\n".join(f"Paragraph number {i} text." for i in range(5)), encoding="utf-8")
    seen = []
    run(translate_txt_file(str(src), str(tmp_path / "o.txt"), make_client(), progress_cb=lambda d, t, s, x: seen.append((d, t))))
    assert seen[-1] == (5, 5)

    ev = asyncio.Event()
    ev.set()
    res = run(translate_txt_file(str(src), str(tmp_path / "o2.txt"), make_client(), cancel_event=ev))
    assert res["status"] == "cancelled" and not (tmp_path / "o2.txt").exists()


def test_txt_missing_input(tmp_path, make_client):
    with pytest.raises(FileNotFoundError):
        run(translate_txt_file(str(tmp_path / "nope.txt"), str(tmp_path / "o.txt"), make_client()))


def test_txt_many_batches_order_preserved(tmp_path, make_client):
    src = tmp_path / "in.txt"
    paras = [f"Unique sentence number {i}." for i in range(40)]
    src.write_text("\n\n".join(paras), encoding="utf-8")
    cl = make_client(batch_items=4, batch_chars=800)
    run(translate_txt_file(str(src), str(tmp_path / "o.txt"), cl))
    out = (tmp_path / "o.txt").read_text(encoding="utf-8").split("\n\n")
    assert out == [f"VI({p})" for p in paras]


# ------------------------------------------------------------------ DOCX
def _make_docx(path):
    d = docx.Document()
    d.add_heading("Chapter One", level=1)
    p = d.add_paragraph()
    r = p.add_run("Bold lead. ")
    r.bold = True
    p.add_run("Plain tail of the paragraph.")
    d.add_paragraph("1234")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Cell alpha text"
    t.rows[0].cells[1].text = "Cell beta text"
    d.save(str(path))


def test_docx_end_to_end(tmp_path, make_client):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _make_docx(src)
    res = run(translate_docx_file(str(src), str(dst), make_client()))
    assert res["status"] == "completed" and res["total_items"] == 4
    out = docx.Document(str(dst))
    texts = [p.text for p in out.paragraphs]
    assert "VI(Chapter One)" in texts
    assert any(t.startswith("VI(Bold lead.") for t in texts)
    assert "1234" in texts  # numbers not translated
    cells = [c.text for c in out.tables[0].rows[0].cells]
    assert cells == ["VI(Cell alpha text)", "VI(Cell beta text)"]


def test_docx_bilingual_keeps_source(tmp_path, make_client):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _make_docx(src)
    run(translate_docx_file(str(src), str(dst), make_client(), bilingual=True))
    alltext = "\n".join(p.text for p in docx.Document(str(dst)).paragraphs)
    assert "Chapter One" in alltext and "VI(Chapter One)" in alltext


def test_docx_cancel(tmp_path, make_client):
    src = tmp_path / "in.docx"
    _make_docx(src)
    ev = asyncio.Event()
    ev.set()
    assert run(translate_docx_file(str(src), str(tmp_path / "o.docx"), make_client(), cancel_event=ev))["status"] == "cancelled"


# ------------------------------------------------------------------ EPUB
CONTAINER = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""
OPF = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>T</dc:title><dc:identifier id="id">x</dc:identifier></metadata>
<manifest>
<item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
<item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>
<item id="img" href="cover.png" media-type="image/png"/>
</manifest>
<spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""
CH = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{t}</title></head>
<body><h1>{t} heading</h1>
<p>{t} first paragraph text.</p>
<p>{t} second paragraph with <em>emphasis</em> inside.</p>
<p>42</p>
<pre>code stays here</pre></body></html>"""
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


@pytest.fixture
def tiny_epub(tmp_path):
    p = tmp_path / "book.epub"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", CONTAINER)
        z.writestr("OEBPS/content.opf", OPF)
        z.writestr("OEBPS/ch1.xhtml", CH.format(t="Alpha"))
        z.writestr("OEBPS/ch2.xhtml", CH.format(t="Beta"))
        z.writestr("OEBPS/cover.png", PNG)
    return p


def test_epub_end_to_end(tmp_path, tiny_epub, make_client):
    dst = tmp_path / "out.epub"
    cl = make_client()
    res = run(translate_epub_file(str(tiny_epub), str(dst), cl))
    assert res["status"] == "completed" and res["total_items"] >= 6
    with zipfile.ZipFile(dst) as z:
        assert z.namelist()[0] == "mimetype"
        assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        assert z.read("OEBPS/cover.png") == PNG
        assert z.read("OEBPS/content.opf").decode() == OPF
        ch1 = z.read("OEBPS/ch1.xhtml").decode("utf-8")
        ch2 = z.read("OEBPS/ch2.xhtml").decode("utf-8")
    assert "VI(" in ch1 and "VI(" in ch2
    assert "Alpha first paragraph text." in ch1  # inside VI(...)
    assert "code stays here" in ch1 and "VI(code" not in ch1
    assert "<p>42</p>" in ch1


def test_epub_bilingual_has_both(tmp_path, tiny_epub, make_client):
    dst = tmp_path / "out.epub"
    run(translate_epub_file(str(tiny_epub), str(dst), make_client(), bilingual=True))
    with zipfile.ZipFile(dst) as z:
        ch1 = z.read("OEBPS/ch1.xhtml").decode("utf-8")
    assert "Alpha first paragraph text." in ch1 and "VI(Alpha first paragraph text.)" in ch1


def test_epub_cancel_and_missing(tmp_path, tiny_epub, make_client):
    ev = asyncio.Event()
    ev.set()
    res = run(translate_epub_file(str(tiny_epub), str(tmp_path / "o.epub"), make_client(), cancel_event=ev))
    assert res["status"] == "cancelled"
    with pytest.raises(FileNotFoundError):
        run(translate_epub_file(str(tmp_path / "x.epub"), str(tmp_path / "o.epub"), make_client()))


def test_epub_chapter_limit_translates_only_first_chapter(tmp_path, tiny_epub, make_client):
    dst = tmp_path / "out.epub"
    cl = make_client(chapter_limit=1)
    run(translate_epub_file(str(tiny_epub), str(dst), cl))
    with zipfile.ZipFile(dst) as z:
        ch1 = z.read("OEBPS/ch1.xhtml").decode()
        ch2 = z.read("OEBPS/ch2.xhtml").decode()
    assert "VI(" in ch1 and "VI(" not in ch2
